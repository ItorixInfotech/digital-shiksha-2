"""Admin: category-wise cutoff Excel template download + upload (CET Cell figures copied into a simple sheet)."""
import csv
import io
import re
from typing import Dict, List, Optional, Tuple

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from lib.db import db
from models.content import CutoffImportResult, ImportIssue
from routers.admin import require_admin
from routers.predictor import EXAMS, fmt

router = APIRouter()

HEADERS = ["College", "Course/Branch", "Exam", "Category", "Closing value"]
CATEGORIES = ["General", "OBC", "EWS", "SC", "ST"]
MAX_BYTES = 5 * 1024 * 1024


def normalize_category(raw: str) -> Optional[str]:
    """Accepts General/Open/GOPEN(S), OBC/OBC-NCL/GOBC(S), EWS, SC/GSC(S), ST/GST(S)."""
    s = re.sub(r"[^A-Z]", "", (raw or "").upper())
    if s.startswith("L"):  # CET Cell ladies-quota codes (LOPENS ...) are not supported
        return None
    s = re.sub(r"^G", "", s) if re.match(r"^G(OPEN|OBC|SC|ST)", s) else s
    s = re.sub(r"(NCL)$", "", s)
    s = re.sub(r"[SHO]$", "", s) if s not in ("OBC", "EWS") else s
    return {"GENERAL": "General", "GEN": "General", "OPEN": "General", "GN": "General", "UR": "General",
            "OBC": "OBC", "EWS": "EWS", "SC": "SC", "ST": "ST"}.get(s)


def _wb_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/admin/cutoffs/template")
async def cutoff_template(_: str = Depends(require_admin)):
    colleges = await db.colleges.find({}, {"_id": 0, "name": 1, "slug": 1, "city": 1, "cutoffs": 1}).sort("name", 1).to_list(1000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cutoffs"
    ws.append(HEADERS)
    for c in colleges:
        for r in c.get("cutoffs", []):
            if r.get("exam") in EXAMS and r.get("value") is not None:
                ws.append([c["name"], r["branch"], r["exam"], r.get("category") or "General", r["value"]])
    head_fill = PatternFill("solid", fgColor="0A2540")
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
    for col, w in zip("ABCDE", (58, 36, 16, 12, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    dv_exam = DataValidation(type="list", formula1='"' + ",".join(EXAMS) + '"', allow_blank=False)
    dv_cat = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=False)
    ws.add_data_validation(dv_exam)
    ws.add_data_validation(dv_cat)
    dv_exam.add("C2:C5000")
    dv_cat.add("D2:D5000")

    cl = wb.create_sheet("Colleges")
    cl.append(["College (use this exact name or the ID)", "College ID", "City"])
    for c in colleges:
        cl.append([c["name"], c["slug"], c["city"]])
    for cell in cl[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
    cl.column_dimensions["A"].width, cl.column_dimensions["B"].width = 62, 26

    ins = wb.create_sheet("Instructions")
    for line in [
        "Digital Shiksha - category-wise cutoff upload",
        "",
        "1. Fill one row per College + Course/Branch + Exam + Category on the 'Cutoffs' sheet.",
        "2. College: exact name or College ID from the 'Colleges' sheet.",
        "3. Exam: " + ", ".join(EXAMS) + ".",
        "4. Category: General, OBC, EWS, SC, ST (CET Cell codes like GOPENS, GOBCS, GSCS, GSTS, EWS also work; ladies quota L* is not supported).",
        "5. Closing value: MHT CET / JEE Main / MAH MBA CET = percentile (e.g. 98.75); NEET UG = score out of 720; JEE Advanced = closing rank (category rank for reserved categories).",
        "6. Upload in Admin -> Cutoffs. Use 'Validate' first; existing rows with the same College + Course + Exam + Category are updated, new ones are added.",
        "7. Rows you delete from this sheet are NOT removed from the website - edit the college in Admin -> Colleges to remove a cutoff.",
    ]:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=14, color="DC2626")
    ins.column_dimensions["A"].width = 140
    wb.move_sheet("Instructions", offset=-2)
    wb.active = 1
    return Response(content=_wb_bytes(wb), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="digital-shiksha-cutoff-template.xlsx"'})


def _read_rows(name: str, data: bytes) -> List[Tuple[int, list]]:
    if name.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [(i, row) for i, row in enumerate(csv.reader(io.StringIO(text)), 1)]
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the file. Upload the .xlsx template (or a .csv).")
    ws = wb["Cutoffs"] if "Cutoffs" in wb.sheetnames else wb.worksheets[0]
    return [(i, list(row)) for i, row in enumerate(ws.iter_rows(values_only=True), 1)]


@router.post("/admin/cutoffs/upload", response_model=CutoffImportResult)
async def upload_cutoffs(file: UploadFile = File(...), dry_run: bool = Query(True), _: str = Depends(require_admin)):
    if not (file.filename or "").lower().endswith((".xlsx", ".csv")):
        raise HTTPException(status_code=400, detail="Please upload an .xlsx or .csv file")
    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="File too large (max 5 MB)")
    rows = _read_rows(file.filename or "", data)
    if not rows:
        raise HTTPException(status_code=400, detail="The file is empty")
    header = [str(h or "").strip().lower() for h in rows[0][1]]
    want = [h.lower() for h in HEADERS]
    if header[:5] != want:
        raise HTTPException(status_code=400, detail=f"First row must be: {', '.join(HEADERS)}")

    colleges = await db.colleges.find({}, {"_id": 0}).to_list(1000)
    lookup: Dict[str, dict] = {}
    for c in colleges:
        for key in (c["slug"], c["name"], c.get("short_name") or ""):
            if key:
                lookup[key.strip().lower()] = c

    issues: List[ImportIssue] = []
    changes: Dict[str, Dict[Tuple[str, str, str], dict]] = {}
    read = 0
    for idx, row in rows[1:]:
        cells = (list(row) + [None] * 5)[:5]
        if all(v is None or str(v).strip() == "" for v in cells):
            continue
        read += 1
        college_raw, branch, exam_raw, cat_raw, value_raw = cells
        college = lookup.get(str(college_raw or "").strip().lower())
        if not college:
            issues.append(ImportIssue(row=idx, message=f"Unknown college '{college_raw}' - use the name or ID from the Colleges sheet"))
            continue
        branch = str(branch or "").strip()
        if not branch:
            issues.append(ImportIssue(row=idx, message="Course/Branch is empty"))
            continue
        exam = next((e for e in EXAMS.values() if e.name.lower() == str(exam_raw or "").strip().lower()), None)
        if not exam:
            issues.append(ImportIssue(row=idx, message=f"Unsupported exam '{exam_raw}' (use: {', '.join(EXAMS)})"))
            continue
        category = normalize_category(str(cat_raw or ""))
        if not category:
            issues.append(ImportIssue(row=idx, message=f"Unsupported category '{cat_raw}' (use: {', '.join(CATEGORIES)})"))
            continue
        try:
            value = float(str(value_raw).replace(",", "").strip())
        except (TypeError, ValueError):
            issues.append(ImportIssue(row=idx, message=f"Closing value '{value_raw}' is not a number"))
            continue
        if not (exam.min <= value <= exam.max):
            issues.append(ImportIssue(row=idx, message=f"Closing value {value:g} is outside {exam.min:g}-{exam.max:g} for {exam.name}"))
            continue
        key = (branch.lower(), exam.name, category)
        changes.setdefault(college["slug"], {})[key] = {
            "exam": exam.name, "branch": branch, "category": category, "value": value, "cutoff": fmt(exam.metric, value),
        }

    added = updated = 0
    by_slug = {c["slug"]: c for c in colleges}
    for slug, rows_for in changes.items():
        existing = by_slug[slug].get("cutoffs", [])
        index = {((r.get("branch") or "").lower(), r.get("exam"), r.get("category") or "General"): i for i, r in enumerate(existing)}
        merged = list(existing)
        for key, new in rows_for.items():
            if key in index:
                merged[index[key]] = new
                updated += 1
            else:
                merged.append(new)
                added += 1
        if not dry_run:
            await db.colleges.update_one({"slug": slug}, {"$set": {"cutoffs": merged}})
    return CutoffImportResult(dry_run=dry_run, rows_read=read, added=added, updated=updated, skipped=len(issues),
                              colleges_affected=len(changes), issues=issues[:200])
